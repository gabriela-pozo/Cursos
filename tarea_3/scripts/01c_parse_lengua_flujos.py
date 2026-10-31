"""Extrae de los tabulados INEI (Educación) dos tablas por distrito que la tubería principal no usaba:
 - cuadro 2: lengua materna de la población de 3 y más años -> lengua_distrito.csv
 - cuadro 8: ubicación de la institución educativa (mismo distrito, otro distrito, otro país) por grupo de edad -> flujo_distrito.csv
Solo conteos. Uso: python 01c_parse_lengua_flujos.py <carpeta inei_tabulados> <carpeta processed>"""
import openpyxl, re, glob, csv, sys, os
SRC, OUT = sys.argv[1], sys.argv[2]
AGE = re.compile(r'^(Menores de 1|De \d+ a \d+|\d+ años?|\d+ y más|De \d+ y más)', re.I)
def num(x):
    try: return int(str(x).replace(',', '').strip())
    except Exception: return None
def walk(ws, minrow=5):
    pad = lambda r: tuple(r) + (None,) * (14 - len(r))
    dept = prov = dist = None
    for r in ws.iter_rows(min_row=minrow, values_only=True):
        r = pad(r); a = r[0]
        if a is None or str(a).strip() == '': continue
        a = str(a).strip()
        if a.startswith('PROVINCIA '): prov = a[10:]; dist = None; yield ('prov', dept, prov, None, None, r); continue
        if a in ('URBANA', 'RURAL'): yield ('row', dept, prov, dist, a, r); continue
        if a.startswith('DISTRITO '): dist = a[9:]; yield ('dist', dept, prov, dist, None, r); continue
        if a.isupper() and not AGE.match(a) and dept is None: dept = a; yield ('dept', dept, None, None, None, r); continue
        if a.isupper() and not AGE.match(a) and a not in ('HOMBRE', 'MUJER'): dept = a; prov = dist = None; continue
        yield ('row', dept, prov, dist, a, r)
def hoja(wb, n):
    for w in wb.worksheets:
        c = next(w.iter_rows(min_row=1, max_row=1, values_only=True), [None])[0]
        if c and re.match(r'^CUADRO N.{0,3}\s*%d\s*:' % n, str(c).strip()): return w
INDIG = ('Quechua', 'Aimara', 'Ashaninka', 'Awajún', 'Shipibo', 'Otros idiomas o lenguas indí')
len_rows, flu_rows = [], []
for f in sorted(glob.glob(SRC + '/*_Educacion.xlsx')):
    code = os.path.basename(f)[:2]
    if code == '00': continue
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    # --- cuadro 2: lengua materna (bloque "total", antes de Hombre/Mujer/URBANA/RURAL)
    ws = hoja(wb, 2); cur = None; fin = False; acc = None
    for kind, dept, prov, dist, lab, r in walk(ws):
        if kind == 'dist':
            if acc: len_rows.append(acc)
            cur = (dept, prov, dist); fin = False; acc = {'k': (code, *cur), 'total': num(r[1]) or 0, 'quechua': 0, 'aimara': 0, 'amaz': 0, 'cast': 0, 'ind': 0}; continue
        if kind in ('prov', 'dept'):
            if acc: len_rows.append(acc); acc = None
            cur = None; continue
        if kind == 'row' and cur and acc:
            if lab in ('Hombre', 'Mujer', 'URBANA', 'RURAL'): fin = True; continue
            if fin: continue
            v = num(r[1]) or 0
            if lab.startswith('Quechua'): acc['quechua'] = v; acc['ind'] += v
            elif lab.startswith('Aimara'): acc['aimara'] = v; acc['ind'] += v
            elif lab.startswith(INDIG[2:]) or lab.startswith('Awaj') or lab.startswith('Shipibo') or lab.startswith('Otros idiomas o lenguas ind'): acc['amaz'] += v; acc['ind'] += v
            elif lab.startswith('Castellano'): acc['cast'] = v
    if acc: len_rows.append(acc); acc = None
    # --- cuadro 8: ubicación de la IE
    ws = hoja(wb, 8); cur = None; fin = False
    for kind, dept, prov, dist, lab, r in walk(ws):
        if kind == 'dist': cur = (dept, prov, dist); fin = False; continue
        if kind in ('prov', 'dept'): cur = None; continue
        if kind == 'row' and cur:
            if lab in ('Hombre', 'Mujer', 'URBANA', 'RURAL'): fin = True; continue
            if fin: continue
            if lab.startswith('De '): flu_rows.append([code, *cur, lab, num(r[1]), num(r[2]), num(r[3]), num(r[4])])
with open(OUT + '/lengua_distrito.csv', 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh); w.writerow(['cod_ambito', 'dept', 'prov', 'dist', 'total_3mas', 'quechua', 'aimara', 'otras_indigenas', 'castellano', 'indigena_total'])
    for a in len_rows: w.writerow([*a['k'], a['total'], a['quechua'], a['aimara'], a['amaz'], a['cast'], a['ind']])
with open(OUT + '/flujo_distrito.csv', 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh); w.writerow(['cod_ambito', 'dept', 'prov', 'dist', 'grupo', 'total', 'en_este_distrito', 'en_otro_distrito', 'en_otro_pais']); w.writerows(flu_rows)
print('lengua distritos', len(len_rows), '| flujo filas', len(flu_rows), 'distritos', len({(r[0], r[2], r[3]) for r in flu_rows}))
