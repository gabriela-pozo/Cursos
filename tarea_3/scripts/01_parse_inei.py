import openpyxl, re, glob, csv, sys, os
SRC=sys.argv[1]; OUT=sys.argv[2]
AGE=re.compile(r'^(Menores de 1|De \d+ a \d+|\d+ años?|\d+ y más|De \d+ y más)',re.I)
def num(x):
    try: return int(str(x).replace(',','').strip())
    except: return None
def walk(ws, minrow=5):
    pad=lambda r:tuple(r)+(None,)*(12-len(r))
    dept=prov=dist=None
    for r in ws.iter_rows(min_row=minrow,values_only=True):
        r=pad(r); a=r[0]
        if a is None or str(a).strip()=='' : continue
        a=str(a).strip()
        if a.startswith('PROVINCIA '): prov=a[10:]; dist=None; level='prov'; yield ('prov',dept,prov,None,None,r); continue
        if a in ('URBANA','RURAL'): yield ('row',dept,prov,dist,a,r); continue
        if a.startswith('DISTRITO '): dist=a[9:]; yield ('dist',dept,prov,dist,None,r); continue
        if a.isupper() and not AGE.match(a) and dept is None: dept=a; yield ('dept',dept,None,None,None,r); continue
        if a.isupper() and not AGE.match(a) and not a.startswith(('PROVINCIA','DISTRITO')) and a not in ('HOMBRE','MUJER'):
            # new department block inside same sheet (unlikely)
            dept=a; prov=dist=None; continue
        yield ('row',dept,prov,dist,a,r)

# Cuadro 2 Indicadores: edades simples por distrito
rows=[]
for f in sorted(glob.glob(SRC+'/*_Indicadores_demograficos.xlsx')):
    code=os.path.basename(f)[:2]
    if code=='00': continue
    ws=openpyxl.load_workbook(f,read_only=True,data_only=True)['CUADRO N°2']
    for kind,dept,prov,dist,lab,r in walk(ws):
        if kind=='row' and dist and lab:
            rows.append([code,dept,prov,dist,lab,num(r[1]),num(r[4]),num(r[7])])
with open(OUT+'/edades_distrito.csv','w',newline='',encoding='utf-8') as fh:
    w=csv.writer(fh); w.writerow(['cod_ambito','dept','prov','dist','edad','total','urbano','rural']); w.writerows(rows)
print('edades filas',len(rows), 'distritos', len({(r[0],r[2],r[3]) for r in rows}))

# Educación cuadro 7: asistencia 3-24 por grupo normativo, distrito, Total sexo (primer bloque antes de Hombre)
rows=[]
for f in sorted(glob.glob(SRC+'/*_Educacion.xlsx')):
    code=os.path.basename(f)[:2]
    if code=='00': continue
    wb=openpyxl.load_workbook(f,read_only=True,data_only=True)
    ws=None
    for w_ in wb.worksheets:
        c=next(w_.iter_rows(min_row=1,max_row=1,values_only=True),[None])[0]
        if c and re.match(r'^CUADRO N.{0,3}\s*7\s*:',str(c).strip()): ws=w_; break
    if ws is None: print('sin cuadro 7',f); continue
    cur=None; seen_sex=False
    for kind,dept,prov,dist,lab,r in walk(ws):
        if kind=='dist': cur=(dept,prov,dist); seen_sex=False; continue
        if kind in('prov','dept'): cur=None; continue
        if kind=='row' and cur:
            if lab in('Hombre','Mujer','URBANA','RURAL'): seen_sex=True; continue
            if seen_sex: continue
            if lab.startswith('De '): rows.append([code,*cur,lab,num(r[1]),num(r[2]),num(r[3])])
with open(OUT+'/asistencia_distrito.csv','w',newline='',encoding='utf-8') as fh:
    w=csv.writer(fh); w.writerow(['cod_ambito','dept','prov','dist','grupo','total','asiste','no_asiste']); w.writerows(rows)
print('asistencia filas',len(rows), 'distritos', len({(r[0],r[2],r[3]) for r in rows}))
