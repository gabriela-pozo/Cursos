"""Extrae tablas adicionales de los tabulados INEI 2025, por distrito:
 - Indicadores demográficos, cuadro 1: población total, censada y omitida  -> omision_distrito.csv
 - Educación, cuadro 10: uso de dispositivos e internet por grupo de edad  -> internet_distrito.csv
Uso: python 01b_parse_extra.py <carpeta_tabulados> <salida_dir>"""
import openpyxl, re, glob, csv, sys, os
SRC, OUT = sys.argv[1], sys.argv[2]
AGE = re.compile(r'^(Menores de 1|De \d+ a \d+|\d+ años?|\d+ y más|De \d+ y más)', re.I)
def num(x):
    try: return int(str(x).replace(',', '').strip())
    except Exception: return None
def find_sheet(wb, n):
    for w in wb.worksheets:
        c = next(w.iter_rows(min_row=1, max_row=1, values_only=True), [None])[0]
        if c and re.match(rf'^CUADRO N.{{0,3}}\s*{n}\s*:', str(c).strip()): return w
    raise KeyError(f'cuadro {n}')
def walk(ws, minrow=5):
    dept = prov = dist = None
    for r in ws.iter_rows(min_row=minrow, values_only=True):
        r = tuple(r) + (None,) * (12 - len(r)); a = r[0]
        if a is None or str(a).strip() == '': continue
        a = str(a).strip()
        if a.startswith('PROVINCIA '): prov = a[10:]; dist = None; yield ('prov', dept, prov, None, a, r); continue
        if a in ('URBANA', 'RURAL', 'HOMBRE', 'MUJER'): yield ('row', dept, prov, dist, a, r); continue
        if a.startswith('DISTRITO '): dist = a[9:]; yield ('dist', dept, prov, dist, a, r); continue
        if a.isupper() and not AGE.match(a) and dept is None: dept = a; yield ('dept', dept, None, None, a, r); continue
        if a.isupper() and not AGE.match(a): dept = a; prov = dist = None; continue
        yield ('row', dept, prov, dist, a, r)
om, net = [], []
for f in sorted(glob.glob(SRC + '/*_Indicadores_demograficos.xlsx')):
    code = os.path.basename(f)[:2]
    if code == '00': continue
    ws = find_sheet(openpyxl.load_workbook(f, read_only=True, data_only=True), 1)
    for kind, dept, prov, dist, lab, r in walk(ws, 6):
        if kind == 'dist': om.append([code, dept, prov, dist, num(r[1]), num(r[4]), num(r[7])])
for f in sorted(glob.glob(SRC + '/*_Educacion.xlsx')):
    code = os.path.basename(f)[:2]
    if code == '00': continue
    ws = find_sheet(openpyxl.load_workbook(f, read_only=True, data_only=True), 10)
    cur = None; stop = False
    for kind, dept, prov, dist, lab, r in walk(ws):
        if kind == 'dist': cur = (dept, prov, dist); stop = False; continue
        if kind in ('prov', 'dept'): cur = None; continue
        if kind == 'row' and cur:
            if lab in ('HOMBRE', 'MUJER', 'URBANA', 'RURAL', 'Hombre', 'Mujer'): stop = True; continue
            if stop: continue
            if lab.startswith('De '): net.append([code, *cur, lab, num(r[1]), num(r[2]), num(r[3]), num(r[4]), num(r[5]), num(r[6])])
with open(OUT + '/omision_distrito.csv', 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh); w.writerow(['cod_ambito', 'dept', 'prov', 'dist', 'pob_total_ajustada', 'pob_censada', 'pob_omitida']); w.writerows(om)
with open(OUT + '/internet_distrito.csv', 'w', newline='', encoding='utf-8') as fh:
    w = csv.writer(fh); w.writerow(['cod_ambito', 'dept', 'prov', 'dist', 'grupo', 'total', 'computadora', 'tableta', 'celular_apps', 'internet', 'ninguno']); w.writerows(net)
print('omisión: distritos', len(om), '| total', sum(r[4] for r in om), 'censada', sum(r[5] for r in om), 'omitida', sum(r[6] for r in om))
print('internet: filas', len(net), 'distritos', len({(r[0], r[2], r[3]) for r in net}))
